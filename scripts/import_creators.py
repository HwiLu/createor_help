"""通用博主真实数据导入脚本（多平台）。

支持从 JSON / CSV / Excel(.xlsx) 导入博主数据，覆盖：
  抖音 / 小红书 / B站 / 微信公众号 / 微信视频号

字段（中英文表头均可）：
  平台(platform)        必填：抖音/小红书/B站/微信公众号/微信视频号，或 DOUYIN/XIAOHONGSHU/BILIBILI/WECHAT_OFFICIAL_ACCOUNT/WECHAT_CHANNEL
  板块(category)        必填：美妆/健身/生活/足球/工作/旅游/游戏，或 BEAUTY/FITNESS/LIFESTYLE/FOOTBALL/WORK/TRAVEL/GAMING
  账号名(username)      必填：用于搜索的唯一账号名
  昵称(nickname)        可选
  粉丝数(follower_count) 可选，纯数字或带单位（如 1200万 / 1.2亿 / 12,000,000）
  头像(avatar_url)      可选
  主页(profile_url)     可选，不填则按平台自动生成搜索链接

导入策略（幂等）：
  按「平台 + 账号名」去重：已存在则更新（昵称/粉丝数/链接），不存在则新增。
  导入完成后，按「平台 + 板块」内粉丝数降序自动重算排名。

用法（在项目根目录）：
  python scripts/import_creators.py --json data.json
  python scripts/import_creators.py --csv  data.csv
  python scripts/import_creators.py --xlsx data.xlsx
  python scripts/import_creators.py --csv data.csv --dry-run            # 只预览不写库
  python scripts/import_creators.py --csv data.csv --clear-all          # 先清空全部再导入
  python scripts/import_creators.py --csv data.csv --clear-platform 小红书 抖音
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
import urllib.parse
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.db.base import Base  # noqa: E402
from app.db.session import SessionLocal, engine  # noqa: E402
from app.models import TopCreator, follower_tier  # noqa: E402
from app.models.enums import (  # noqa: E402
    CREATOR_CATEGORY_LABELS,
    PLATFORM_LABELS,
    CreatorCategory,
    Platform,
)

# ---- 中英文标签 <-> 枚举值 映射 ----
PLATFORM_LABEL_TO_VALUE = {label: p.value for p, label in PLATFORM_LABELS.items()}
CATEGORY_LABEL_TO_VALUE = {label: c.value for c, label in CREATOR_CATEGORY_LABELS.items()}
PLATFORM_VALUES = {p.value for p in Platform}
CATEGORY_VALUES = {c.value for c in CreatorCategory}

# ---- 字段别名（CSV 表头/JSON 键均支持中英文）----
FIELD_ALIASES = {
    "platform": ["platform", "平台"],
    "category": ["category", "板块", "类目", "分类"],
    "username": ["username", "账号名", "用户名", "账号"],
    "nickname": ["nickname", "昵称", "名称"],
    "follower_count": ["follower_count", "粉丝数", "粉丝", "粉丝数量"],
    "avatar_url": ["avatar_url", "头像", "头像链接", "头像url"],
    "profile_url": ["profile_url", "主页", "主页链接", "主页url", "链接"],
    "rank": ["rank", "排名"],
}


def parse_platform(value: str) -> str | None:
    """返回平台枚举值；无法识别返回 None。"""
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    if s in PLATFORM_LABEL_TO_VALUE:
        return PLATFORM_LABEL_TO_VALUE[s]
    up = s.upper()
    if up in PLATFORM_VALUES:
        return up
    return None


def parse_category(value: str) -> str | None:
    """返回板块枚举值；无法识别返回 None。"""
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    if s in CATEGORY_LABEL_TO_VALUE:
        return CATEGORY_LABEL_TO_VALUE[s]
    up = s.upper()
    if up in CATEGORY_VALUES:
        return up
    return None


def parse_follower_count(value) -> int | None:
    """解析粉丝数：支持 12000000 / 1200万 / 1.2亿 / 12,000,000 / 1.2w / 3.5k。失败返回 None。"""
    if value is None:
        return 0
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    s = str(value).strip().replace(",", "").replace("，", "").replace(" ", "")
    if not s:
        return 0
    units = [
        ("亿", 100_000_000),
        ("w", 10_000),
        ("W", 10_000),
        ("万", 10_000),
        ("k", 1_000),
        ("K", 1_000),
    ]
    for suffix, mul in units:
        if s.endswith(suffix):
            try:
                return int(float(s[: -len(suffix)]) * mul)
            except ValueError:
                return None
    try:
        return int(float(s))
    except ValueError:
        return None


def profile_url(platform_value: str, username: str) -> str:
    kw = urllib.parse.quote(username)
    mapping = {
        Platform.XIAOHONGSHU.value: f"https://www.xiaohongshu.com/search_result?keyword={kw}",
        Platform.DOUYIN.value: f"https://www.douyin.com/search/{kw}",
        Platform.BILIBILI.value: f"https://search.bilibili.com/all?keyword={kw}",
        Platform.WECHAT_OFFICIAL_ACCOUNT.value: f"https://weixin.sogou.com/weixin?type=1&query={kw}",
        Platform.WECHAT_CHANNEL.value: f"https://weixin.sogou.com/weixin?type=2&query={kw}",
    }
    return mapping.get(platform_value, f"https://www.baidu.com/s?wd={kw}")


def normalize_row(raw: dict) -> dict:
    """把中英文/别名键归一化为标准字段名。"""
    norm: dict = {}
    for std_key, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in raw:
                v = raw[alias]
                if v is None:
                    norm[std_key] = None
                else:
                    s = str(v).strip()
                    norm[std_key] = s if s else None
                break
    return norm


def load_rows(json_path=None, csv_path=None, xlsx_path=None) -> list[dict]:
    rows: list[dict] = []
    if json_path:
        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        items = data if isinstance(data, list) else data.get("creators", [])
        rows.extend(items)
    if csv_path:
        with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
            rows.extend(list(csv.DictReader(f)))
    if xlsx_path:
        try:
            import openpyxl  # type: ignore
        except ImportError:
            print("[错误] 读取 .xlsx 需要 openpyxl，请先 `pip install openpyxl`，")
            print("       或改用 CSV（Excel 中「另存为」→ CSV UTF-8）再导入。")
            sys.exit(1)
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows.append(dict(zip(headers, row)))
    return rows


def build_parsed(rows: list[dict]) -> tuple[list[dict], list[str]]:
    """解析并校验所有行，返回 (有效记录列表, 错误信息列表)。"""
    parsed: list[dict] = []
    errors: list[str] = []
    for idx, raw in enumerate(rows, start=1):
        row = normalize_row(raw)
        line_no = idx + 1 if True else idx  # CSV 首行是表头
        platform_value = parse_platform(row.get("platform"))
        category_value = parse_category(row.get("category"))
        username = row.get("username")

        problems = []
        if platform_value is None:
            problems.append(f"平台无法识别：{row.get('platform')!r}")
        if category_value is None:
            problems.append(f"板块无法识别：{row.get('category')!r}")
        if not username:
            problems.append("缺少账号名(username)")

        follower_count = parse_follower_count(row.get("follower_count"))
        if follower_count is None:
            problems.append(f"粉丝数无法解析：{row.get('follower_count')!r}")

        if problems:
            errors.append(f"第 {line_no} 行：{'；'.join(problems)} —— {raw!r}")
            continue

        parsed.append(
            {
                "platform": platform_value,
                "category": category_value,
                "username": username,
                "nickname": row.get("nickname") or None,
                "follower_count": follower_count or 0,
                "follower_display": row.get("follower_display") if "follower_display" in row else None,
                "avatar_url": row.get("avatar_url") or None,
                "profile_url": row.get("profile_url") or profile_url(platform_value, username),
                "rank": row.get("rank"),
            }
        )
    return parsed, errors


def preview(parsed: list[dict]) -> None:
    """按「平台 + 板块 + 粉丝数降序」打印预览排名。"""
    grouped = defaultdict(list)
    for p in parsed:
        grouped[(p["platform"], p["category"])].append(p)
    print(f"\n预览（共 {len(parsed)} 条）：")
    for (platform_value, category_value), items in sorted(grouped.items()):
        items_sorted = sorted(items, key=lambda x: x["follower_count"], reverse=True)
        label = PLATFORM_LABELS.get(Platform(platform_value), platform_value)
        cat_label = CREATOR_CATEGORY_LABELS.get(CreatorCategory(category_value), category_value)
        print(f"\n[{label} · {cat_label}]")
        for rank, it in enumerate(items_sorted, start=1):
            print(f"  #{rank}  {it['username']:<20} 粉丝 {it['follower_count']:>12,}")


def upsert(parsed: list[dict], clear_all: bool, clear_platforms: list[str]) -> int:
    """写库。返回导入条数。"""
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        if clear_all:
            deleted = db.query(TopCreator).delete()
            print(f"[信息] 已清空全部 {deleted} 条旧数据。")
        for platform_value in clear_platforms:
            deleted = db.query(TopCreator).filter(TopCreator.platform == platform_value).delete()
            label = PLATFORM_LABELS.get(Platform(platform_value), platform_value)
            print(f"[信息] 已清空 {label} 的 {deleted} 条旧数据。")

        imported = 0
        for p in parsed:
            existing = (
                db.query(TopCreator)
                .filter(TopCreator.platform == p["platform"], TopCreator.username == p["username"])
                .first()
            )
            follower_display = p["follower_display"] or follower_tier(p["follower_count"])
            if existing:
                existing.category = p["category"]
                if p["nickname"]:
                    existing.nickname = p["nickname"]
                existing.follower_count = p["follower_count"]
                existing.follower_display = follower_display
                if p["avatar_url"]:
                    existing.avatar_url = p["avatar_url"]
                existing.profile_url = p["profile_url"]
                existing.is_featured = True
                existing.status = "ACTIVE"
            else:
                db.add(
                    TopCreator(
                        platform=p["platform"],
                        category=p["category"],
                        username=p["username"],
                        nickname=p["nickname"],
                        avatar_url=p["avatar_url"],
                        profile_url=p["profile_url"],
                        follower_count=p["follower_count"],
                        follower_display=follower_display,
                        rank=0,
                        is_featured=True,
                        status="ACTIVE",
                    )
                )
            imported += 1

        db.flush()
        # 按「平台 + 板块」内粉丝数降序重算排名
        all_rows = db.query(TopCreator).order_by(
            TopCreator.platform, TopCreator.category, TopCreator.follower_count.desc(), TopCreator.id
        ).all()
        counters: dict[tuple, int] = defaultdict(int)
        for r in all_rows:
            counters[(r.platform, r.category)] += 1
            r.rank = counters[(r.platform, r.category)]

        db.commit()
        return imported
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="通用博主真实数据导入脚本")
    parser.add_argument("--json", help="JSON 数据文件路径")
    parser.add_argument("--csv", help="CSV 数据文件路径")
    parser.add_argument("--xlsx", help="Excel(.xlsx) 数据文件路径")
    parser.add_argument("--dry-run", action="store_true", help="只解析预览，不写库")
    parser.add_argument("--clear-all", action="store_true", help="导入前清空全部旧数据")
    parser.add_argument(
        "--clear-platform",
        nargs="+",
        default=[],
        help="导入前清空指定平台（可用中文或枚举值，如：小红书 抖音）",
    )
    args = parser.parse_args()

    if not any([args.json, args.csv, args.xlsx]):
        parser.print_help()
        sys.exit(1)

    rows = load_rows(json_path=args.json, csv_path=args.csv, xlsx_path=args.xlsx)
    if not rows:
        print("[错误] 没有读到任何数据行，请检查文件内容与表头。")
        sys.exit(1)

    parsed, errors = build_parsed(rows)
    if errors:
        print("[错误] 以下数据行存在问题，已跳过：")
        for e in errors:
            print("  -", e)
    if not parsed:
        print("[错误] 没有可导入的有效数据。")
        sys.exit(1)

    preview(parsed)

    clear_platforms = [parse_platform(x) for x in args.clear_platform]
    invalid_clear = [x for x, v in zip(args.clear_platform, clear_platforms) if v is None]
    if invalid_clear:
        print(f"[错误] --clear-platform 含无法识别的平台：{invalid_clear}")
        sys.exit(1)

    if args.dry_run:
        print("\n[提示] --dry-run 模式：以上为预览，未写入数据库。")
        return

    imported = upsert(parsed, args.clear_all, clear_platforms)
    print(f"\n[OK] 已导入 {imported} 条博主数据。")


if __name__ == "__main__":
    main()
